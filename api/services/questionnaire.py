"""Questionnaire processing service.

Handles:
- Text extraction from uploaded documents (PDF, Excel, Word, CSV)
- AI-powered question extraction from unstructured text
- AI-powered draft answer generation using company data
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
import os
import re
from typing import Any

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from api.models import (
    Company,
    EmissionReport,
    Questionnaire,
    QuestionnaireQuestion,
)

logger = logging.getLogger(__name__)

# LLM model identifiers — override via env vars to upgrade models without code changes
_ANTHROPIC_MODEL: str = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-20250514")
_OPENAI_MODEL: str = os.getenv("OPENAI_MODEL", "gpt-4o-mini")

# LLM call timeout in seconds — prevents hanging on unresponsive API endpoints
_LLM_TIMEOUT_SECONDS: float = float(os.getenv("LLM_TIMEOUT_SECONDS", "30"))


# ── Document text extraction ─────────────────────────────────────────


def extract_text_from_pdf(content: bytes) -> str:
    """Extract text from a PDF file."""
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = []
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            return "\n\n".join(pages)
    except ImportError:
        logger.warning("pdfplumber not installed; PDF extraction disabled")
        return ""


def extract_text_from_docx(content: bytes) -> str:
    """Extract text from a Word document."""
    try:
        import docx

        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except ImportError:
        logger.warning("python-docx not installed; DOCX extraction disabled")
        return ""


def extract_text_from_xlsx(content: bytes) -> str:
    """Extract text from an Excel file."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except ImportError:
        logger.warning("openpyxl not installed; XLSX extraction disabled")
        return ""


def extract_text_from_csv(content: bytes) -> str:
    """Extract text from a CSV file."""
    text = content.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader)


EXTRACTORS = {
    "pdf": extract_text_from_pdf,
    "docx": extract_text_from_docx,
    "xlsx": extract_text_from_xlsx,
    "csv": extract_text_from_csv,
}


def extract_text(content: bytes, file_type: str) -> str:
    """Route to the correct extractor based on file type."""
    extractor = EXTRACTORS.get(file_type)
    if not extractor:
        return content.decode("utf-8", errors="replace")
    return extractor(content)


# ── Question extraction (rule-based fallback) ────────────────────────

# Common sustainability questionnaire patterns
_QUESTION_PATTERNS = [
    re.compile(r"^\s*(?:Q\d+|Question\s+\d+)[.:]\s*(.+)", re.IGNORECASE),
    re.compile(r"^\s*\d+[.)]\s+(.+\?)\s*$"),
    re.compile(r"^\s*[-•]\s+(.+\?)\s*$"),
    re.compile(r"^(.{20,}?\?)\s*$"),
]

_CATEGORY_KEYWORDS = {
    "emissions": ["emission", "ghg", "co2", "carbon", "scope 1", "scope 2", "scope 3", "greenhouse"],
    "energy": ["energy", "electricity", "renewable", "solar", "wind", "power", "fuel"],
    "waste": ["waste", "recycl", "circular", "landfill", "disposal"],
    "water": ["water", "consumption", "discharge", "effluent"],
    "transport": ["transport", "logistics", "fleet", "vehicle", "freight", "shipping"],
    "governance": ["governance", "board", "policy", "strategy", "target", "commitment"],
    "supply_chain": ["supply chain", "supplier", "procurement", "sourcing", "vendor"],
    "reporting": ["report", "disclosure", "cdp", "tcfd", "gri", "csrd", "sbti"],
}


def _classify_question(text: str) -> str | None:
    """Classify a question into a sustainability category."""
    lower = text.lower()
    for category, keywords in _CATEGORY_KEYWORDS.items():
        if any(kw in lower for kw in keywords):
            return category
    return None


def extract_questions_rule_based(text: str) -> list[dict[str, Any]]:
    """Extract questions from text using regex patterns."""
    questions: list[dict[str, Any]] = []
    seen = set()

    for line in text.split("\n"):
        line = line.strip()
        if not line or len(line) < 15:
            continue

        for pattern in _QUESTION_PATTERNS:
            match = pattern.match(line)
            if match:
                q_text = match.group(1).strip() if match.lastindex else line.strip()
                if q_text not in seen and len(q_text) > 10:
                    seen.add(q_text)
                    questions.append({
                        "question_text": q_text,
                        "category": _classify_question(q_text),
                    })
                break

    # If no pattern-based questions found, split by sentences ending with ?
    if not questions:
        for sentence in re.split(r"[.!]\s+", text):
            sentence = sentence.strip()
            if sentence.endswith("?") and len(sentence) > 15 and sentence not in seen:
                seen.add(sentence)
                questions.append({
                    "question_text": sentence,
                    "category": _classify_question(sentence),
                })

    return questions


# ── LLM-based extraction ────────────────────────────────────────────

_EXTRACT_QUESTIONS_PROMPT = """You are an expert sustainability questionnaire analyzer. Extract all individual questions from this document text.

For each question, provide:
1. The exact question text
2. A category: emissions, energy, waste, water, transport, governance, supply_chain, or reporting

Return a JSON array of objects with keys: "question_text", "category"

Document text:
{text}

JSON array:"""

_DRAFT_ANSWER_PROMPT = """You are a sustainability reporting expert helping a company respond to a questionnaire.

Company: {company_name}
Industry: {industry}
Region: {region}
Latest Emissions (if available):
- Scope 1: {scope1} tCO2e
- Scope 2: {scope2} tCO2e
- Scope 3: {scope3} tCO2e
- Total: {total} tCO2e

Question: {question}

Write a professional, accurate draft answer (2-4 sentences). If you don't have enough data, note what information would be needed. Be factual, not speculative."""


def _get_llm_client():
    """Lazy-init LLM client."""
    api_key = os.getenv("OPENAI_API_KEY") or os.getenv("ANTHROPIC_API_KEY")
    provider = "anthropic" if os.getenv("ANTHROPIC_API_KEY") else "openai"
    if not api_key:
        return None, provider
    if provider == "anthropic":
        try:
            import anthropic
            return anthropic.Anthropic(api_key=api_key), provider
        except ImportError:
            return None, provider
    else:
        try:
            import openai
            return openai.OpenAI(api_key=api_key), provider
        except ImportError:
            return None, provider


def _llm_call_sync(client: Any, provider: str, prompt: str, max_tokens: int = 2048) -> str:
    """Make a synchronous LLM call (runs in thread pool via caller)."""
    if provider == "anthropic":
        resp = client.messages.create(
            model=_ANTHROPIC_MODEL,
            max_tokens=max_tokens,
            messages=[{"role": "user", "content": prompt}],
        )
        return resp.content[0].text
    else:
        resp = client.chat.completions.create(
            model=_OPENAI_MODEL,
            messages=[{"role": "user", "content": prompt}],
            max_tokens=max_tokens,
            temperature=0.2,
        )
        return resp.choices[0].message.content


async def extract_questions_llm(text: str) -> list[dict[str, Any]]:
    """Extract questions using LLM."""
    client, provider = _get_llm_client()
    if not client:
        return extract_questions_rule_based(text)

    try:
        prompt = _EXTRACT_QUESTIONS_PROMPT.format(text=text[:8000])
        content = await asyncio.wait_for(
            asyncio.to_thread(_llm_call_sync, client, provider, prompt),
            timeout=_LLM_TIMEOUT_SECONDS,
        )

        # Parse JSON array from response
        json_match = re.search(r"\[.*\]", content, re.DOTALL)
        if json_match:
            return json.loads(json_match.group())
        return json.loads(content)
    except (json.JSONDecodeError, KeyError, TypeError, OSError, asyncio.TimeoutError) as e:
        logger.warning("LLM question extraction failed, using rule-based: %s", e)
        return extract_questions_rule_based(text)


async def generate_draft_answer(
    question: str,
    company_name: str,
    industry: str,
    region: str,
    scope1: float = 0,
    scope2: float = 0,
    scope3: float = 0,
    total: float = 0,
) -> tuple[str, float]:
    """Generate a draft answer for a question. Returns (answer, confidence)."""
    client, provider = _get_llm_client()

    if not client:
        return _draft_answer_rule_based(question, company_name, industry), 0.3

    try:
        prompt = _DRAFT_ANSWER_PROMPT.format(
            company_name=company_name,
            industry=industry,
            region=region,
            scope1=scope1,
            scope2=scope2,
            scope3=scope3,
            total=total,
            question=question,
        )
        answer = await asyncio.wait_for(
            asyncio.to_thread(_llm_call_sync, client, provider, prompt, 512),
            timeout=_LLM_TIMEOUT_SECONDS,
        )
        return answer.strip(), 0.75
    except (KeyError, TypeError, OSError, asyncio.TimeoutError) as e:
        logger.warning("LLM draft answer failed: %s", e)
        return _draft_answer_rule_based(question, company_name, industry), 0.3


def _draft_answer_rule_based(question: str, company_name: str, industry: str) -> str:
    """Generate a basic template answer without LLM."""
    category = _classify_question(question)
    templates = {
        "emissions": f"{company_name} tracks greenhouse gas emissions across Scope 1, 2, and 3 in accordance with the GHG Protocol Corporate Standard. Detailed emissions data is available in our annual sustainability report.",
        "energy": f"{company_name} monitors energy consumption across all facilities and is actively pursuing energy efficiency improvements and renewable energy procurement.",
        "waste": f"{company_name} has implemented waste reduction programs and tracks waste generation, recycling rates, and disposal methods across operations.",
        "transport": f"{company_name} monitors transportation and logistics emissions, including fleet operations and freight distribution, and evaluates opportunities for optimization.",
        "governance": f"{company_name} maintains a sustainability governance structure with board-level oversight. Our sustainability strategy is integrated into overall business planning.",
        "supply_chain": f"{company_name} engages with suppliers on sustainability performance and tracks supply chain emissions as part of our Scope 3 inventory.",
        "reporting": f"{company_name} discloses environmental performance through established reporting frameworks. We are committed to transparent and accurate sustainability reporting.",
    }
    return templates.get(category, f"{company_name} is committed to addressing this area as part of our sustainability strategy. Further details can be provided upon request.")


# ── Orchestration ────────────────────────────────────────────────────


async def process_questionnaire(
    db: AsyncSession,
    questionnaire_id: str,
    company_id: str,
) -> Questionnaire:
    """Extract questions from a questionnaire and generate draft answers."""
    result = await db.execute(
        select(Questionnaire).where(
            Questionnaire.id == questionnaire_id,
            Questionnaire.company_id == company_id,
            Questionnaire.deleted_at.is_(None),
        )
    )
    questionnaire = result.scalar_one_or_none()
    if not questionnaire:
        raise ValueError("Questionnaire not found")

    # Update status
    questionnaire.status = "extracting"
    await db.commit()

    # Extract questions
    text = questionnaire.extracted_text or ""
    raw_questions = await extract_questions_llm(text)

    if not raw_questions:
        raw_questions = extract_questions_rule_based(text)

    # Delete existing questions to prevent duplicates on re-extraction
    await db.execute(
        delete(QuestionnaireQuestion).where(
            QuestionnaireQuestion.questionnaire_id == questionnaire_id
        )
    )

    # Get company info for draft answers
    company_result = await db.execute(
        select(Company).where(Company.id == company_id, Company.deleted_at.is_(None))
    )
    company = company_result.scalar_one()

    # Get latest emission report if available
    report_result = await db.execute(
        select(EmissionReport)
        .where(
            EmissionReport.company_id == company_id,
            EmissionReport.deleted_at.is_(None),
        )
        .order_by(EmissionReport.created_at.desc())
        .limit(1)
    )
    report = report_result.scalar_one_or_none()

    scope1 = report.scope1 if report else 0
    scope2 = report.scope2 if report else 0
    scope3 = report.scope3 if report else 0
    total = report.total if report else 0

    # Create question records + draft answers
    for i, q in enumerate(raw_questions, start=1):
        draft, confidence = await generate_draft_answer(
            question=q["question_text"],
            company_name=company.name,
            industry=company.industry,
            region=company.region,
            scope1=scope1,
            scope2=scope2,
            scope3=scope3,
            total=total,
        )
        db.add(QuestionnaireQuestion(
            questionnaire_id=questionnaire_id,
            question_number=i,
            question_text=q["question_text"],
            category=q.get("category"),
            ai_draft_answer=draft,
            confidence=confidence,
        ))

    questionnaire.status = "extracted"
    await db.commit()
    await db.refresh(questionnaire)
    return questionnaire
